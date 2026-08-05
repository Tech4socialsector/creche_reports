import frappe
import frappe
import json
from openpyxl import Workbook
from openpyxl.styles import Protection
from io import BytesIO

# @frappe.whitelist(allow_guest=True)
# def get_all_budget_items():

#     return frappe.get_all(
#         "Budget and Expense items list",
#         fields=[
#             "name",
#             "budget_main_head",
#             "budget_sub_head",
#             "type_of_expenses"
#         ],
#         order_by="name asc"
#     )


@frappe.whitelist()
def get_all_budget_items():

    # --------------------------------------------------
    # GET LOGGED IN USER
    # --------------------------------------------------

    logged_in_user = frappe.session.user

    # --------------------------------------------------
    # GET MAIN HEAD SEQUENCE MAPPING
    # --------------------------------------------------

    main_head_sequence = {
        d.name: (d.sequence_id or 0)
        for d in frappe.get_all(
            "Budget main head",
            fields=["name", "sequence_id"]
        )
    }

    # --------------------------------------------------
    # SORT HELPER
    # --------------------------------------------------

    def sort_items(items):
        return sorted(
            items,
            key=lambda d: (
                main_head_sequence.get(d.get("budget_main_head"), 0),
                d.get("budget_sub_head") or "",
                d.get("name") or ""
            )
        )

    # --------------------------------------------------
    # FIND IMPORT TEMPLATE SETTINGS
    # --------------------------------------------------

    template_settings_name = frappe.db.get_value(
        "Import Template settings",
        {
            "user": logged_in_user
        }
    )

    # --------------------------------------------------
    # IF NO TEMPLATE SETTINGS FOUND
    # RETURN ALL BUDGET ITEMS SORTED
    # --------------------------------------------------

    if not template_settings_name:

        return sort_items(
            frappe.get_all(
                "Budget and Expense items list",
                fields=[
                    "name",
                    "budget_main_head",
                    "budget_sub_head",
                    "type_of_expenses"
                ]
            )
        )

    # --------------------------------------------------
    # GET DOCUMENT
    # --------------------------------------------------

    template_settings = frappe.get_doc(
        "Import Template settings",
        template_settings_name
    )

    # --------------------------------------------------
    # BUILD BUDGET ITEMS FROM CHILD TABLE
    # --------------------------------------------------

    budget_items = []

    for row in template_settings.table_jndv:

        budget_items.append({

            "name": row.type_of_expenses_id,

            "budget_main_head": row.budget_main_head,

            "budget_sub_head": row.budget_sub_head,

            "type_of_expenses": row.type_of_expenses
        })

    # --------------------------------------------------
    # IF CHILD TABLE IS EMPTY
    # RETURN ALL BUDGET ITEMS SORTED
    # --------------------------------------------------

    if not budget_items:

        return sort_items(
            frappe.get_all(
                "Budget and Expense items list",
                fields=[
                    "name",
                    "budget_main_head",
                    "budget_sub_head",
                    "type_of_expenses"
                ]
            )
        )

    # --------------------------------------------------
    # RETURN CHILD TABLE ITEMS SORTED
    # --------------------------------------------------

    return sort_items(budget_items)


# ==================================================== Budget Import Template ===============================================================================
# @frappe.whitelist(allow_guest=True)
# def download_budget_template(data=None):

#     if isinstance(data, str):
#         data = json.loads(data)

#     budget_items = frappe.get_all(
#         "Budget and Expense items list",
#         fields=[
#             "name",
#             "budget_main_head",
#             "budget_sub_head",
#             "type_of_expenses"
#         ],
#         order_by="name asc"
#     )

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Creche Budget"

#     headers = [
#         "Budget reference name",
#         "Partner ID",
#         "Partner Name",
#         "Grant ID",
#         "No of creches",
#         "State",
#         "Financial year",
#         "Date of approval",
#         "Start Date",
#         "End Date",

#         "Type of expenses ID (Budget Items List)",
#         "Budget main head (Budget Items List)",
#         "Budget sub head (Budget Items List)",
#         "Type of expenses (Budget Items List)",
#         "Year 1 (Budget Items List)",
#         "Year 2 (Budget Items List)",
#         "Year 3 (Budget Items List)",
#         "Total Amount (Budget Items List)",
#         "Notes (Budget Items List)",

#         "Total budget"
#     ]

#     ws.append(headers)

#     first_row = True
#     start_child_row = 2
#     excel_row = 2

#     for item in budget_items:

#         row = []

#         if first_row:

#             row.extend([
#                 data.get("budget_reference_name"),
#                 data.get("partner_id"),
#                 data.get("partner_name"),
#                 data.get("grant_id"),
#                 data.get("no_of_creches"),
#                 data.get("state"),
#                 data.get("financial_year"),
#                 data.get("date_of_approval"),
#                 data.get("start_date"),
#                 data.get("end_date")
#             ])

#             first_row = False

#         else:
#             row.extend([""] * 10)

#         # Total Amount = Year1 + Year2 + Year3
#         total_amount_formula = f'=SUM(O{excel_row}:Q{excel_row})'

#         row.extend([
#             item.name,
#             item.budget_main_head,
#             item.budget_sub_head,
#             item.type_of_expenses,

#             "",  # Year 1
#             "",  # Year 2
#             "",  # Year 3

#             total_amount_formula,
#             ""   # Notes
#         ])

#         total_budget_formula = (
#             f'=SUM(R{start_child_row}:R{start_child_row + len(budget_items)-1})'
#         )

#         if excel_row == 2:
#             row.append(total_budget_formula)
#         else:
#             row.append("")

#         ws.append(row)

#         excel_row += 1

#     # Unlock all cells
#     for row_cells in ws.iter_rows():
#         for cell in row_cells:
#             cell.protection = Protection(locked=False)

#     # Lock header
#     for cell in ws[1]:
#         cell.protection = Protection(locked=True)

#     # Lock parent fields
#     parent_columns = [
#         "A", "B", "C", "D", "E",
#         "F", "G", "H", "I", "J"
#     ]

#     for col in parent_columns:
#         for row_no in range(2, excel_row):
#             ws[f"{col}{row_no}"].protection = Protection(locked=True)

#     # Lock child master fields
#     child_master_columns = ["K", "L", "M", "N"]

#     for col in child_master_columns:
#         for row_no in range(2, excel_row):
#             ws[f"{col}{row_no}"].protection = Protection(locked=True)

#     # Lock formula columns
#     formula_columns = ["R", "T"]

#     for col in formula_columns:
#         for row_no in range(2, excel_row):
#             ws[f"{col}{row_no}"].protection = Protection(locked=True)

#     # Editable:
#     # O = Year 1
#     # P = Year 2
#     # Q = Year 3
#     # S = Notes

#     ws.protection.sheet = True
#     ws.protection.password = "1234"

#     for column_cells in ws.columns:

#         max_length = max(
#             len(str(cell.value or ""))
#             for cell in column_cells
#         )

#         column_letter = column_cells[0].column_letter

#         ws.column_dimensions[column_letter].width = max_length + 5

#     file_stream = BytesIO()
#     wb.save(file_stream)

#     frappe.response["filename"] = "creche_budget_template.xlsx"
#     frappe.response["filecontent"] = file_stream.getvalue()
#     frappe.response["type"] = "binary"



# @frappe.whitelist(allow_guest=True)
# def download_budget_template(data=None):

#     import json

#     from io import BytesIO

#     from openpyxl import Workbook
#     from openpyxl.styles import (
#         Font,
#         PatternFill
#     )

#     if isinstance(data, str):
#         data = json.loads(data)

#     data = data or {}

#     # --------------------------------------------------
#     # GET MAIN HEAD SEQUENCE
#     # --------------------------------------------------

#     main_head_sequence = {
#         d.name: d.sequence_id or 0
#         for d in frappe.get_all(
#             "Budget main head",
#             fields=["name", "sequence_id"]
#         )
#     }

#     # --------------------------------------------------
#     # GET ITEMS
#     # --------------------------------------------------

#     budget_items = frappe.get_all(
#         "Budget and Expense items list",
#         fields=[
#             "name",
#             "budget_main_head",
#             "budget_sub_head",
#             "type_of_expenses"
#         ]
#     )

#     # --------------------------------------------------
#     # SORT ITEMS
#     # --------------------------------------------------

#     budget_items = sorted(
#         budget_items,
#         key=lambda d: (
#             main_head_sequence.get(
#                 d.budget_main_head, 0
#             ),
#             d.budget_sub_head or "",
#             d.name or ""
#         )
#     )

#     # --------------------------------------------------
#     # CREATE WORKBOOK
#     # --------------------------------------------------

#     wb = Workbook()

#     ws = wb.active
#     ws.title = "Creche Budget"

#     # --------------------------------------------------
#     # HEADERS
#     # --------------------------------------------------

#     headers = [

#         "Budget reference name",
#         "Partner ID",
#         "Partner Name",
#         "Grant ID",
#         "No of creches",
#         "State",
#         "Financial year",
#         "Date of approval",
#         "Start Date",
#         "End Date",

#         "Type of expenses ID (Budget Items List)",
#         "Budget main head (Budget Items List)",
#         "Budget sub head (Budget Items List)",
#         "Type of expenses (Budget Items List)",

#         "Year 1 (Budget Items List)",
#         "Year 2 (Budget Items List)",
#         "Year 3 (Budget Items List)",

#         "Total Amount (Budget Items List)",

#         "Notes (Budget Items List)",

#         "Total budget"
#     ]

#     ws.append(headers)

#     # --------------------------------------------------
#     # HEADER STYLE
#     # --------------------------------------------------

#     header_fill = PatternFill(
#         start_color="D9EAF7",
#         end_color="D9EAF7",
#         fill_type="solid"
#     )

#     for cell in ws[1]:

#         cell.font = Font(bold=True)
#         cell.fill = header_fill

#     # --------------------------------------------------
#     # ADD ROWS
#     # --------------------------------------------------

#     for idx, item in enumerate(
#         budget_items,
#         start=2
#     ):

#         row = [

#             # Parent fields only first row
#             data.get("budget_reference_name")
#             if idx == 2 else "",

#             data.get("partner_id")
#             if idx == 2 else "",

#             data.get("partner_name")
#             if idx == 2 else "",

#             data.get("grant_id")
#             if idx == 2 else "",

#             data.get("no_of_creches")
#             if idx == 2 else "",

#             data.get("state")
#             if idx == 2 else "",

#             data.get("financial_year")
#             if idx == 2 else "",

#             data.get("date_of_approval")
#             if idx == 2 else "",

#             data.get("start_date")
#             if idx == 2 else "",

#             data.get("end_date")
#             if idx == 2 else "",

#             # Child table fields
#             item.name,
#             item.budget_main_head,
#             item.budget_sub_head,
#             item.type_of_expenses,

#             # Editable currency fields
#             0.0,
#             0.0,
#             0.0,

#             # Total Amount
#             0.0,

#             # Notes
#             "",

#             # Total Budget
#             0.0 if idx == 2 else ""
#         ]

#         ws.append(row)

#     # --------------------------------------------------
#     # TOTAL COLUMN STYLE
#     # --------------------------------------------------

#     total_fill = PatternFill(
#         start_color="E8F4EA",
#         end_color="E8F4EA",
#         fill_type="solid"
#     )

#     currency_columns = [
#         "O", "P", "Q", "R", "T"
#     ]

#     for col in currency_columns:

#         for row_no in range(
#             2,
#             ws.max_row + 1
#         ):

#             cell = ws[f"{col}{row_no}"]

#             cell.fill = total_fill
#             cell.number_format = '#,##0.00'

#     # --------------------------------------------------
#     # FILTERS + FREEZE
#     # --------------------------------------------------

#     ws.auto_filter.ref = ws.dimensions
#     ws.freeze_panes = "A2"

#     # --------------------------------------------------
#     # AUTO WIDTH
#     # --------------------------------------------------

#     for column in ws.columns:

#         length = max(
#             len(str(cell.value or ""))
#             for cell in column
#         )

#         ws.column_dimensions[
#             column[0].column_letter
#         ].width = min(length + 5, 50)

#     # --------------------------------------------------
#     # HIDE COLUMNS
#     # --------------------------------------------------

#     hidden_columns = [

#         "B",  # Partner ID
#         "K",  # Expense Item ID
#         "S",  # Notes
#         "T"   # Total Budget
#     ]

#     for col in hidden_columns:

#         ws.column_dimensions[
#             col
#         ].hidden = True

#     # --------------------------------------------------
#     # FILE NAME
#     # --------------------------------------------------

#     filename = frappe.scrub(
#         data.get(
#             "budget_reference_name"
#         ) or "budget"
#     ).replace("-", "_")

#     # --------------------------------------------------
#     # DOWNLOAD
#     # --------------------------------------------------

#     output = BytesIO()

#     wb.save(output)

#     frappe.response["filename"] = (
#         f"{filename}_creche_budget_template.xlsx"
#     )

#     frappe.response["filecontent"] = (
#         output.getvalue()
#     )

#     frappe.response["type"] = "binary"


# @frappe.whitelist(allow_guest=True)
# def download_budget_template(data=None):

#     import json
#     from io import BytesIO
#     from openpyxl import Workbook
#     from openpyxl.styles import Font, PatternFill

#     if isinstance(data, str):
#         data = json.loads(data)

#     # --------------------------------------------------
#     # GET MAIN HEAD SEQUENCE
#     # --------------------------------------------------

#     main_head_sequence = {
#         d.name: d.sequence_id or 0
#         for d in frappe.get_all(
#             "Budget main head",
#             fields=["name", "sequence_id"]
#         )
#     }

#     # --------------------------------------------------
#     # GET ITEMS
#     # --------------------------------------------------

#     budget_items = frappe.get_all(
#         "Budget and Expense items list",
#         fields=[
#             "name",
#             "budget_main_head",
#             "budget_sub_head",
#             "type_of_expenses"
#         ]
#     )

#     # --------------------------------------------------
#     # SORT ITEMS
#     # --------------------------------------------------

#     budget_items = sorted(
#         budget_items,
#         key=lambda d: (
#             main_head_sequence.get(d.budget_main_head, 0),
#             d.budget_sub_head or "",
#             d.name or ""
#         )
#     )

#     # --------------------------------------------------
#     # CREATE WORKBOOK
#     # --------------------------------------------------

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Creche Budget"

#     headers = [
#        "Budget reference name",
#         "Partner ID",
#         "Partner Name",
#         "Grant ID",
#         "No of creches",
#         "State",
#         "Financial year",
#         "Date of approval",
#         "Start Date",
#         "End Date",

#         "Type of expenses ID (Budget Items List)",
#         "Budget main head (Budget Items List)",
#         "Budget sub head (Budget Items List)",
#         "Type of expenses (Budget Items List)",
#         "Year 1 (Budget Items List)",
#         "Year 2 (Budget Items List)",
#         "Year 3 (Budget Items List)",
#         "Total Amount (Budget Items List)",
#         "Notes (Budget Items List)",
#         "Total budget"
#     ]

#     ws.append(headers)

#     # --------------------------------------------------
#     # HEADER STYLE
#     # --------------------------------------------------

#     header_fill = PatternFill(
#         start_color="D9EAF7",
#         end_color="D9EAF7",
#         fill_type="solid"
#     )

#     for cell in ws[1]:
#         cell.font = Font(bold=True)
#         cell.fill = header_fill

#     # --------------------------------------------------
#     # ADD ROWS
#     # --------------------------------------------------

#     start_row = 2

#     for idx, item in enumerate(budget_items, start=2):

#         row = [

#             data.get("budget_reference_name") if idx == 2 else "",
#             data.get("partner_id") if idx == 2 else "",
#             data.get("partner_name") if idx == 2 else "",
#             data.get("grant_id") if idx == 2 else "",
#             data.get("no_of_creches") if idx == 2 else "",
#             data.get("state") if idx == 2 else "",
#             data.get("financial_year") if idx == 2 else "",
#             data.get("date_of_approval") if idx == 2 else "",
#             data.get("start_date") if idx == 2 else "",
#             data.get("end_date") if idx == 2 else "",

#             item.name,
#             item.budget_main_head,
#             item.budget_sub_head,
#             item.type_of_expenses,

#             0,  # Year 1
#             0,  # Year 2
#             0,  # Year 3

#             f'=SUM(O{idx}:Q{idx})',

#             "",  # Notes

#             f'=SUM(R{start_row}:R{len(budget_items)+1})'
#             if idx == 2 else 0
#         ]

#         ws.append(row)

#     # --------------------------------------------------
#     # TOTAL AMOUNT COLUMN COLOR
#     # --------------------------------------------------

#     total_fill = PatternFill(
#         start_color="E8F4EA",
#         end_color="E8F4EA",
#         fill_type="solid"
#     )

#     for row in range(2, ws.max_row + 1):
#         ws[f"R{row}"].fill = total_fill

#     # --------------------------------------------------
#     # FILTERS + FREEZE
#     # --------------------------------------------------

#     ws.auto_filter.ref = ws.dimensions
#     ws.freeze_panes = "A2"

#     # --------------------------------------------------
#     # AUTO WIDTH
#     # --------------------------------------------------

#     for column in ws.columns:

#         length = max(
#             len(str(cell.value or ""))
#             for cell in column
#         )

#         ws.column_dimensions[
#             column[0].column_letter
#         ].width = min(length + 5, 50)

#     # --------------------------------------------------
#     # HIDE COLUMNS
#     # --------------------------------------------------

#     hidden_columns = [
#         "B",  # Partner ID
#         "K",  # Type of expenses ID
#         "S",  # Notes
#         "T"   # Total budget
#     ]

#     for col in hidden_columns:
#         ws.column_dimensions[col].hidden = True

#     # --------------------------------------------------
#     # FILE NAME
#     # --------------------------------------------------

#     filename = frappe.scrub(
#         data.get("budget_reference_name") or "budget"
#     ).replace("-", "_")

#     # --------------------------------------------------
#     # DOWNLOAD
#     # --------------------------------------------------

#     output = BytesIO()
#     wb.save(output)

#     frappe.response["filename"] = (
#         f"{filename}_creche_budget_template.xlsx"
#     )

#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "binary"



@frappe.whitelist()
def download_budget_template(data=None):

    import json
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    # --------------------------------------------------
    # PARSE DATA
    # --------------------------------------------------

    if isinstance(data, str):
        data = json.loads(data)

    data = data or {}

    # --------------------------------------------------
    # GET LOGGED IN USER
    # --------------------------------------------------

    logged_in_user = frappe.session.user

    # --------------------------------------------------
    # GET BUDGET ITEMS
    # Prefer the user's own "Import Template settings" (per-user override)
    # if one exists with items configured; otherwise fall back to the
    # shared master list so every user can generate a template, matching
    # the same fallback used by get_all_budget_items().
    # --------------------------------------------------

    budget_items = []

    template_settings_name = frappe.db.get_value(
        "Import Template settings",
        {
            "user": logged_in_user
        }
    )

    if template_settings_name:

        template_settings = frappe.get_doc(
            "Import Template settings",
            template_settings_name
        )

        for row in template_settings.table_jndv:

            budget_items.append({

                "name": row.type_of_expenses_id,

                "budget_main_head": row.budget_main_head,

                "budget_sub_head": row.budget_sub_head,

                "type_of_expenses": row.type_of_expenses
            })

    if not budget_items:

        main_head_sequence = {
            d.name: (d.sequence_id or 0)
            for d in frappe.get_all(
                "Budget main head",
                fields=["name", "sequence_id"]
            )
        }

        budget_items = sorted(
            frappe.get_all(
                "Budget and Expense items list",
                fields=[
                    "name",
                    "budget_main_head",
                    "budget_sub_head",
                    "type_of_expenses"
                ]
            ),
            key=lambda d: (
                main_head_sequence.get(d.get("budget_main_head"), 0),
                d.get("budget_sub_head") or "",
                d.get("name") or ""
            )
        )

    # --------------------------------------------------
    # VALIDATION
    # --------------------------------------------------

    if not budget_items:

        frappe.throw(
            "No Budget and Expense items are configured. Please add items to "
            "'Budget and Expense items list' before generating a template."
        )

    # --------------------------------------------------
    # CREATE WORKBOOK
    # --------------------------------------------------

    wb = Workbook()

    ws = wb.active

    ws.title = "Creche Budget"

    # --------------------------------------------------
    # HEADERS
    # --------------------------------------------------

    headers = [

        "Budget reference name",
        "Partner ID",
        "Partner Name",
        "Grant ID",
        "No of creches",
        "Block",
        "District",
        "State",
        "Financial year",
        "Date of approval",
        "Start Date",
        "End Date",

        "Type of expenses ID (Budget Items List)",
        "Budget main head (Budget Items List)",
        "Budget sub head (Budget Items List)",
        "Type of expenses (Budget Items List)",

        "Year 1 (Budget Items List)",
        "Year 2 (Budget Items List)",
        "Year 3 (Budget Items List)",

        "Total Amount (Budget Items List)",
        "Notes (Budget Items List)",

        "Total budget"
    ]

    ws.append(headers)

    # --------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------

    header_fill = PatternFill(
        start_color="D9EAF7",
        end_color="D9EAF7",
        fill_type="solid"
    )

    for cell in ws[1]:

        cell.font = Font(bold=True)

        cell.fill = header_fill

    # --------------------------------------------------
    # COLUMN INDEXES
    # --------------------------------------------------
    #
    # Q = Year 1
    # R = Year 2
    # S = Year 3
    # T = Total Amount
    # U = Notes
    # V = Total Budget
    #
    # --------------------------------------------------

    total_budget_formula = (
        f"=SUM(T2:T{len(budget_items)+1})"
    )

    # --------------------------------------------------
    # ADD DATA ROWS
    # --------------------------------------------------

    for idx, item in enumerate(budget_items, start=2):

        row = [

            # ------------------------------------------
            # PARENT FIELDS
            # ------------------------------------------

            data.get("budget_reference_name")
            if idx == 2 else "",

            data.get("partner_id")
            if idx == 2 else "",

            data.get("partner_name")
            if idx == 2 else "",

            data.get("grant_id")
            if idx == 2 else "",

            data.get("no_of_creches")
            if idx == 2 else "",

            data.get("block")
            if idx == 2 else "",

            data.get("district")
            if idx == 2 else "",

            data.get("state")
            if idx == 2 else "",

            data.get("financial_year")
            if idx == 2 else "",

            data.get("date_of_approval")
            if idx == 2 else "",

            data.get("start_date")
            if idx == 2 else "",

            data.get("end_date")
            if idx == 2 else "",

            # ------------------------------------------
            # CHILD TABLE VALUES
            # ------------------------------------------

            item.get("name"),

            item.get("budget_main_head"),

            item.get("budget_sub_head"),

            item.get("type_of_expenses"),

            # ------------------------------------------
            # YEAR VALUES
            # ------------------------------------------

            0,  # Year 1
            0,  # Year 2
            0,  # Year 3

            # ------------------------------------------
            # TOTAL AMOUNT
            # ------------------------------------------

            f"=SUM(Q{idx}:S{idx})",

            # ------------------------------------------
            # NOTES
            # ------------------------------------------

            "",

            # ------------------------------------------
            # TOTAL BUDGET
            # ------------------------------------------

            total_budget_formula
            if idx == 2 else ""
        ]

        ws.append(row)

    # --------------------------------------------------
    # STYLE TOTAL AMOUNT COLUMN
    # --------------------------------------------------

    total_fill = PatternFill(
        start_color="E8F4EA",
        end_color="E8F4EA",
        fill_type="solid"
    )

    for row in range(2, ws.max_row + 1):

        ws[f"T{row}"].fill = total_fill

    # --------------------------------------------------
    # STYLE TOTAL BUDGET COLUMN
    # --------------------------------------------------

    budget_fill = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    for row in range(2, ws.max_row + 1):

        ws[f"V{row}"].fill = budget_fill

    # --------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------

    ws.freeze_panes = "A2"

    # --------------------------------------------------
    # ENABLE FILTERS
    # --------------------------------------------------

    ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------
    # AUTO WIDTH
    # --------------------------------------------------

    for column in ws.columns:

        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        adjusted_width = min(
            max_length + 5,
            50
        )

        ws.column_dimensions[
            column[0].column_letter
        ].width = adjusted_width

    # --------------------------------------------------
    # HIDE COLUMNS
    # --------------------------------------------------

    hidden_columns = [

        "B",  # Partner ID
        "M",  # Type of expenses ID
        "U"   # Notes
    ]

    for col in hidden_columns:

        ws.column_dimensions[col].hidden = True

    # --------------------------------------------------
    # FILE NAME
    # --------------------------------------------------

    filename = frappe.scrub(
        data.get("budget_reference_name")
        or "budget"
    ).replace("-", "_")

    # --------------------------------------------------
    # SAVE FILE
    # --------------------------------------------------

    output = BytesIO()

    wb.save(output)

    frappe.response["filename"] = (
        f"{filename}_creche_budget_template.xlsx"
    )

    frappe.response["filecontent"] = output.getvalue()

    frappe.response["type"] = "binary"

# ==================================================== Utilisation Import Template ===============================================================================

# @frappe.whitelist(allow_guest=True)
# def download_utilisation_template(data=None):

#     import json
#     from io import BytesIO
#     from openpyxl import Workbook
#     from openpyxl.styles import Protection, Font, PatternFill

#     # --------------------------------------------------
#     # PARSE DATA
#     # --------------------------------------------------

#     if isinstance(data, str):
#         data = json.loads(data)

#     data = data or {}

#     # --------------------------------------------------
#     # GET LOGGED IN USER
#     # --------------------------------------------------

#     logged_in_user = frappe.session.user

#     # --------------------------------------------------
#     # FIND IMPORT TEMPLATE SETTINGS
#     # --------------------------------------------------

#     template_settings_name = frappe.db.get_value(
#         "Import Template settings",
#         {
#             "user": logged_in_user
#         }
#     )

#     # --------------------------------------------------
#     # VALIDATION
#     # --------------------------------------------------

#     if not template_settings_name:

#         frappe.throw(
#             f"No Import Template settings found for user: {logged_in_user}"
#         )

#     # --------------------------------------------------
#     # GET DOCUMENT
#     # --------------------------------------------------

#     template_settings = frappe.get_doc(
#         "Import Template settings",
#         template_settings_name
#     )

#     # --------------------------------------------------
#     # GET CHILD TABLE ITEMS
#     # ORDER WILL BE SAME AS CHILD TABLE
#     # --------------------------------------------------

#     items = []

#     for row in template_settings.table_jndv:

#         items.append({

#             "name": row.type_of_expenses_id,

#             "budget_main_head": row.budget_main_head,

#             "budget_sub_head": row.budget_sub_head,

#             "type_of_expenses": row.type_of_expenses
#         })

#     # --------------------------------------------------
#     # VALIDATION
#     # --------------------------------------------------

#     if not items:

#         frappe.throw(
#             "No Import Template Items found for this user"
#         )

#     # --------------------------------------------------
#     # CREATE WORKBOOK
#     # --------------------------------------------------

#     wb = Workbook()

#     ws = wb.active

#     ws.title = "Creche Utilisation"

#     # --------------------------------------------------
#     # HEADERS
#     # --------------------------------------------------

#     headers = [
#         "Budget reference ID",                            # A
#         "Budget reference Name",                          # B
#         "Partner ID",                                     # C
#         "Partner Name",                                   # D
#         "Grant ID",                                       # E
#         "State",                                          # F
#         "No of creches",                                  # G
#         "Month",                                          # H
#         "Financial year",                                 # I
#         "Date",                                           # J

#         "Type of expenses ID (Utilisation Items List)",   # K
#         "Budget main head (Utilisation Items List)",      # L
#         "Budget sub head (Utilisation Items List)",       # M
#         "Type of expenses (Utilisation Items List)",      # N

#         "Total Amount (Utilisation Items List)",          # O
#         "Notes (Utilisation Items List)",                 # P

#         "Total Utilisation",                              # Q
#         "Bank + Cash Balance as at end of month reported",# R
#         "Interest from Bank",                             # S
#         "Declaration"                                     # T
#     ]

#     ws.append(headers)

#     # --------------------------------------------------
#     # HEADER STYLE
#     # --------------------------------------------------

#     header_fill = PatternFill(
#         start_color="D9EAF7",
#         end_color="D9EAF7",
#         fill_type="solid"
#     )

#     for cell in ws[1]:
#         cell.font = Font(bold=True)
#         cell.fill = header_fill

#     # --------------------------------------------------
#     # PARENT DATA
#     # --------------------------------------------------

#     parent_data = [
#         data.get("budget_reference_id"),    # A
#         data.get("budget_reference_name"),  # B
#         data.get("partner_id"),             # C
#         data.get("partner_name"),           # D
#         data.get("grant_id"),               # E
#         data.get("state"),                  # F
#         data.get("no_of_creches"),          # G
#         data.get("month"),                  # H
#         data.get("financial_year"),         # I
#         data.get("date")                    # J
#     ]

#     # --------------------------------------------------
#     # ROW RANGE
#     # --------------------------------------------------

#     start_row = 2
#     end_row = start_row + len(items) - 1

#     # --------------------------------------------------
#     # ADD DATA ROWS
#     # --------------------------------------------------

#     for idx, item in enumerate(items, start=start_row):

#         row = []

#         # Parent data only in first row
#         row.extend(
#             parent_data if idx == start_row else [""] * 10
#         )

#         # Child table values
#         row.extend([
#             item.get("name"),               # K
#             item.get("budget_main_head"),   # L
#             item.get("budget_sub_head"),    # M
#             item.get("type_of_expenses"),   # N
#         ])

#         # Editable Total Amount
#         row.append(0)                       # O

#         # Notes
#         row.append("")                      # P

#         # Final columns only in first row
#         if idx == start_row:
#             row.extend([
#                 f"=SUM(O{start_row}:O{end_row})",  # Q - Total Utilisation
#                 0,                                  # R - Bank + Cash Balance
#                 0,                                  # S - Interest from Bank
#                 0                                   # T - Declaration
#             ])
#         else:
#             row.extend(["", "", "", ""])

#         ws.append(row)

#     # --------------------------------------------------
#     # UNLOCK ALL CELLS FIRST
#     # --------------------------------------------------

#     for row in ws.iter_rows():
#         for cell in row:
#             cell.protection = Protection(locked=False)

#     # --------------------------------------------------
#     # LOCK HEADER ROW
#     # --------------------------------------------------

#     for cell in ws[1]:
#         cell.protection = Protection(locked=True)

#     # --------------------------------------------------
#     # READ-ONLY COLUMNS
#     # Editable: H (Month), I (Financial Year), O (Total Amount), P (Notes)
#     # Everything else in data rows is locked
#     # --------------------------------------------------

#     readonly_columns = [
#         "A",  # Budget reference ID
#         "B",  # Budget reference Name
#         "C",  # Partner ID
#         "D",  # Partner Name
#         "E",  # Grant ID
#         "F",  # State
#         "G",  # No of creches
#         "J",  # Date
#         "K",  # Type of expenses ID
#         "L",  # Budget main head
#         "M",  # Budget sub head
#         "N",  # Type of expenses
#         "Q",  # Total Utilisation
#         "R",  # Bank + Cash Balance
#         "S",  # Interest from Bank
#         "T",  # Declaration
#     ]

#     for col in readonly_columns:
#         for row_no in range(2, end_row + 1):
#             ws[f"{col}{row_no}"].protection = Protection(locked=True)

#     # --------------------------------------------------
#     # STYLE - TOTAL AMOUNT COLUMN (O) - green tint
#     # --------------------------------------------------

#     total_fill = PatternFill(
#         start_color="E8F4EA",
#         end_color="E8F4EA",
#         fill_type="solid"
#     )

#     for row_no in range(2, end_row + 1):
#         ws[f"O{row_no}"].fill = total_fill

#     # --------------------------------------------------
#     # STYLE - TOTAL UTILISATION COLUMN (Q) - yellow tint
#     # --------------------------------------------------

#     budget_fill = PatternFill(
#         start_color="FFF2CC",
#         end_color="FFF2CC",
#         fill_type="solid"
#     )

#     for row_no in range(start_row, end_row + 1):
#         ws[f"Q{row_no}"].fill = budget_fill

#     # --------------------------------------------------
#     # PROTECT SHEET
#     # --------------------------------------------------

#     ws.protection.sheet = True
#     ws.protection.password = "1234"

#     # --------------------------------------------------
#     # FREEZE HEADER
#     # --------------------------------------------------

#     ws.freeze_panes = "A2"

#     # --------------------------------------------------
#     # ENABLE FILTERS
#     # --------------------------------------------------

#     ws.auto_filter.ref = ws.dimensions

#     # --------------------------------------------------
#     # AUTO COLUMN WIDTH
#     # --------------------------------------------------

#     for col in ws.columns:
#         max_length = max(
#             len(str(cell.value or ""))
#             for cell in col
#         )
#         adjusted_width = min(max_length + 5, 50)
#         ws.column_dimensions[col[0].column_letter].width = adjusted_width

#     # --------------------------------------------------
#     # HIDE COLUMNS
#     # --------------------------------------------------

#     hidden_columns = [
#         "C",  # Partner ID
#         "K",  # Type of expenses ID
#     ]

#     for col in hidden_columns:
#         ws.column_dimensions[col].hidden = True

#     # --------------------------------------------------
#     # FILE NAME
#     # --------------------------------------------------

#     month = data.get("month") or "month"

#     financial_year = (
#         str(data.get("financial_year") or "fy")
#         .split("-")[0]  # e.g. "2024-2025" → "2024"
#     )

#     filename = f"{month}_{financial_year}_creche_utilisation_template"

#     # --------------------------------------------------
#     # SAVE FILE
#     # --------------------------------------------------

#     output = BytesIO()

#     wb.save(output)

#     frappe.response["filename"] = f"{filename}.xlsx"
#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "binary"

# @frappe.whitelist(allow_guest=True)
# def download_utilisation_template(data=None):

#     import json
#     from io import BytesIO
#     from openpyxl import Workbook
#     from openpyxl.styles import Font, PatternFill

#     # --------------------------------------------------
#     # PARSE DATA
#     # --------------------------------------------------

#     if isinstance(data, str):
#         data = json.loads(data)

#     data = data or {}

#     # --------------------------------------------------
#     # GET LOGGED IN USER
#     # --------------------------------------------------

#     logged_in_user = frappe.session.user

#     # --------------------------------------------------
#     # FIND IMPORT TEMPLATE SETTINGS
#     # --------------------------------------------------

#     template_settings_name = frappe.db.get_value(
#         "Import Template settings",
#         {
#             "user": logged_in_user
#         }
#     )

#     # --------------------------------------------------
#     # VALIDATION
#     # --------------------------------------------------

#     if not template_settings_name:

#         frappe.throw(
#             f"No Import Template settings found for user: {logged_in_user}"
#         )

#     # --------------------------------------------------
#     # GET DOCUMENT
#     # --------------------------------------------------

#     template_settings = frappe.get_doc(
#         "Import Template settings",
#         template_settings_name
#     )

#     # --------------------------------------------------
#     # GET CHILD TABLE ITEMS
#     # ORDER WILL BE SAME AS CHILD TABLE
#     # --------------------------------------------------

#     items = []

#     for row in template_settings.table_jndv:

#         items.append({

#             "name": row.type_of_expenses_id,

#             "budget_main_head": row.budget_main_head,

#             "budget_sub_head": row.budget_sub_head,

#             "type_of_expenses": row.type_of_expenses
#         })

#     # --------------------------------------------------
#     # VALIDATION
#     # --------------------------------------------------

#     if not items:

#         frappe.throw(
#             "No Import Template Items found for this user"
#         )

#     # --------------------------------------------------
#     # CREATE WORKBOOK
#     # --------------------------------------------------

#     wb = Workbook()

#     ws = wb.active

#     ws.title = "Creche Utilisation"

#     # --------------------------------------------------
#     # HEADERS
#     # --------------------------------------------------

#     headers = [
#         "Budget reference ID",                            # A
#         "Budget reference Name",                          # B
#         "Partner ID",                                     # C
#         "Partner Name",                                   # D
#         "Grant ID",                                       # E
#         "State",                                          # F
#         "No of creches",                                  # G
#         "Month",                                          # H
#         "Financial year",                                 # I
#         "Date",                                           # J

#         "Type of expenses ID (Utilisation Items List)",   # K
#         "Budget main head (Utilisation Items List)",      # L
#         "Budget sub head (Utilisation Items List)",       # M
#         "Type of expenses (Utilisation Items List)",      # N

#         "Total Amount (Utilisation Items List)",          # O
#         "Notes (Utilisation Items List)",                 # P

#         "Total Utilisation",                              # Q
#         "Bank + Cash Balance as at end of month reported",# R
#         "Interest from Bank",                             # S
#         "Declaration"                                     # T
#     ]

#     ws.append(headers)

#     # --------------------------------------------------
#     # HEADER STYLE
#     # --------------------------------------------------

#     header_fill = PatternFill(
#         start_color="D9EAF7",
#         end_color="D9EAF7",
#         fill_type="solid"
#     )

#     for cell in ws[1]:
#         cell.font = Font(bold=True)
#         cell.fill = header_fill

#     # --------------------------------------------------
#     # PARENT DATA
#     # --------------------------------------------------

#     parent_data = [
#         data.get("budget_reference_id"),    # A
#         data.get("budget_reference_name"),  # B
#         data.get("partner_id"),             # C
#         data.get("partner_name"),           # D
#         data.get("grant_id"),               # E
#         data.get("state"),                  # F
#         data.get("no_of_creches"),          # G
#         data.get("month"),                  # H
#         data.get("financial_year"),         # I
#         data.get("date")                    # J
#     ]

#     # --------------------------------------------------
#     # ROW RANGE
#     # --------------------------------------------------

#     start_row = 2
#     end_row = start_row + len(items) - 1

#     # --------------------------------------------------
#     # ADD DATA ROWS
#     # --------------------------------------------------

#     for idx, item in enumerate(items, start=start_row):

#         row = []

#         # Parent data only in first row
#         row.extend(
#             parent_data if idx == start_row else [""] * 10
#         )

#         # Child table values
#         row.extend([
#             item.get("name"),               # K
#             item.get("budget_main_head"),   # L
#             item.get("budget_sub_head"),    # M
#             item.get("type_of_expenses"),   # N
#         ])

#         # Editable Total Amount
#         row.append(0)                       # O

#         # Notes
#         row.append("")                      # P

#         # Final columns only in first row
#         if idx == start_row:
#             row.extend([
#                 f"=SUM(O{start_row}:O{end_row})",  # Q - Total Utilisation
#                 0,                                  # R - Bank + Cash Balance
#                 0,                                  # S - Interest from Bank
#                 0                                   # T - Declaration
#             ])
#         else:
#             row.extend(["", "", "", ""])

#         ws.append(row)

#     # --------------------------------------------------
#     # STYLE - TOTAL AMOUNT COLUMN (O) - green tint
#     # --------------------------------------------------

#     total_fill = PatternFill(
#         start_color="E8F4EA",
#         end_color="E8F4EA",
#         fill_type="solid"
#     )

#     for row_no in range(2, end_row + 1):
#         ws[f"O{row_no}"].fill = total_fill

#     # --------------------------------------------------
#     # STYLE - TOTAL UTILISATION COLUMN (Q) - yellow tint
#     # --------------------------------------------------

#     budget_fill = PatternFill(
#         start_color="FFF2CC",
#         end_color="FFF2CC",
#         fill_type="solid"
#     )

#     for row_no in range(start_row, end_row + 1):
#         ws[f"Q{row_no}"].fill = budget_fill

#     # --------------------------------------------------
#     # FREEZE HEADER
#     # --------------------------------------------------

#     ws.freeze_panes = "A2"

#     # --------------------------------------------------
#     # ENABLE FILTERS
#     # --------------------------------------------------

#     ws.auto_filter.ref = ws.dimensions

#     # --------------------------------------------------
#     # AUTO COLUMN WIDTH
#     # --------------------------------------------------

#     for col in ws.columns:
#         max_length = max(
#             len(str(cell.value or ""))
#             for cell in col
#         )
#         adjusted_width = min(max_length + 5, 50)
#         ws.column_dimensions[col[0].column_letter].width = adjusted_width

#     # --------------------------------------------------
#     # HIDE COLUMNS
#     # --------------------------------------------------

#     hidden_columns = [
#         "A",  # Budget reference ID
#         "C",  # Partner ID
#         "J",  # Date
#         # "K",  # Type of expenses ID
#         "T",  # Declaration
#     ]

#     for col in hidden_columns:
#         ws.column_dimensions[col].hidden = True

#     # --------------------------------------------------
#     # FILE NAME
#     # --------------------------------------------------

#     month = data.get("month") or "month"

#     financial_year = (
#         str(data.get("financial_year") or "fy")
#         .split("-")[0]  # e.g. "2024-2025" → "2024"
#     )

#     filename = f"{month}_{financial_year}_creche_utilisation_template"

#     # --------------------------------------------------
#     # SAVE FILE
#     # --------------------------------------------------

#     output = BytesIO()

#     wb.save(output)

#     frappe.response["filename"] = f"{filename}.xlsx"
#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "binary"



# @frappe.whitelist(allow_guest=True)
# def download_utilisation_template(data=None):

#     import json
#     from io import BytesIO
#     from openpyxl import Workbook
#     from openpyxl.styles import Font, PatternFill

#     # --------------------------------------------------
#     # PARSE DATA
#     # --------------------------------------------------

#     if isinstance(data, str):
#         data = json.loads(data)

#     data = data or {}

#     # --------------------------------------------------
#     # GET LOGGED IN USER
#     # --------------------------------------------------

#     logged_in_user = frappe.session.user

#     # --------------------------------------------------
#     # FIND IMPORT TEMPLATE SETTINGS
#     # --------------------------------------------------

#     template_settings_name = frappe.db.get_value(
#         "Import Template settings",
#         {
#             "user": logged_in_user
#         }
#     )

#     # --------------------------------------------------
#     # GET DOCUMENT
#     # --------------------------------------------------

#     if template_settings_name:

#         template_settings = frappe.get_doc(
#             "Import Template settings",
#             template_settings_name
#         )

#         items = []

#         for row in template_settings.table_jndv:

#             items.append({

#                 "name": row.type_of_expenses_id,

#                 "budget_main_head": row.budget_main_head,

#                 "budget_sub_head": row.budget_sub_head,

#                 "type_of_expenses": row.type_of_expenses
#             })

#     else:

#         items = frappe.get_all(
#             "Budget and Expense items list",
#             fields=[
#                 "name",
#                 "budget_main_head",
#                 "budget_sub_head",
#                 "type_of_expenses"
#             ],
#             order_by="name asc"
#         )

#     # --------------------------------------------------
#     # VALIDATION
#     # --------------------------------------------------

#     if not items:

#         frappe.throw(
#             "No Import Template Items found for this user"
#         )

#     # --------------------------------------------------
#     # CREATE WORKBOOK
#     # --------------------------------------------------

#     wb = Workbook()

#     ws = wb.active

#     ws.title = "Creche Utilisation"

#     # --------------------------------------------------
#     # HEADERS
#     # --------------------------------------------------

#     headers = [
#         "Budget reference ID",                            # A
#         "Budget reference Name",                          # B
#         "Partner ID",                                     # C
#         "Partner Name",                                   # D
#         "Grant ID",                                       # E
#         "Block",                                          # F
#         "District",                                       # G
#         "State",                                          # H
#         "No of creches",                                  # I
#         "Month",                                          # J
#         "Financial year",                                 # K
#         "Date",                                           # L

#         "Type of expenses ID (Utilisation Items List)",   # M
#         "Budget main head (Utilisation Items List)",      # N
#         "Budget sub head (Utilisation Items List)",       # O
#         "Type of expenses (Utilisation Items List)",      # P

#         "Total Amount (Utilisation Items List)",          # Q
#         "Notes (Utilisation Items List)",                 # R

#         "Total Utilisation",                              # S
#         "Bank + Cash Balance as at end of month reported",# T
#         "Interest from Bank",                             # U
#         "Declaration"                                     # V
#     ]

#     ws.append(headers)

#     # --------------------------------------------------
#     # HEADER STYLE
#     # --------------------------------------------------

#     header_fill = PatternFill(
#         start_color="D9EAF7",
#         end_color="D9EAF7",
#         fill_type="solid"
#     )

#     for cell in ws[1]:
#         cell.font = Font(bold=True)
#         cell.fill = header_fill

#     # --------------------------------------------------
#     # PARENT DATA
#     # --------------------------------------------------

#     parent_data = [
#         data.get("budget_reference_id"),    # A
#         data.get("budget_reference_name"),  # B
#         data.get("partner_id"),             # C
#         data.get("partner_name"),           # D
#         data.get("grant_id"),               # E
#         data.get("block"),                  # F
#         data.get("district"),               # G
#         data.get("state"),                  # H
#         data.get("no_of_creches"),          # I
#         data.get("month"),                  # J
#         data.get("financial_year"),         # K
#         data.get("date")                    # L
#     ]

#     # --------------------------------------------------
#     # ROW RANGE
#     # --------------------------------------------------

#     start_row = 2
#     end_row = start_row + len(items) - 1

#     # --------------------------------------------------
#     # ADD DATA ROWS
#     # --------------------------------------------------

#     for idx, item in enumerate(items, start=start_row):

#         row = []

#         # Parent data only in first row
#         row.extend(
#             parent_data if idx == start_row else [""] * 12
#         )

#         # Child table values
#         row.extend([
#             item.get("name"),               # M
#             item.get("budget_main_head"),   # N
#             item.get("budget_sub_head"),    # O
#             item.get("type_of_expenses"),   # P
#         ])

#         # Editable Total Amount
#         row.append(0)                       # Q

#         # Notes
#         row.append("")                      # R

#         # Final columns only in first row
#         if idx == start_row:
#             row.extend([
#                 f"=SUM(Q{start_row}:Q{end_row})",  # S - Total Utilisation
#                 0,                                  # T - Bank + Cash Balance
#                 0,                                  # U - Interest from Bank
#                 1                                   # V - Declaration
#             ])
#         else:
#             row.extend(["", "", "", ""])

#         ws.append(row)

#     # --------------------------------------------------
#     # STYLE - TOTAL AMOUNT COLUMN (Q) - green tint
#     # --------------------------------------------------

#     total_fill = PatternFill(
#         start_color="E8F4EA",
#         end_color="E8F4EA",
#         fill_type="solid"
#     )

#     for row_no in range(2, end_row + 1):
#         ws[f"Q{row_no}"].fill = total_fill

#     # --------------------------------------------------
#     # STYLE - TOTAL UTILISATION COLUMN (S) - yellow tint
#     # --------------------------------------------------

#     budget_fill = PatternFill(
#         start_color="FFF2CC",
#         end_color="FFF2CC",
#         fill_type="solid"
#     )

#     for row_no in range(start_row, end_row + 1):
#         ws[f"S{row_no}"].fill = budget_fill

#     # --------------------------------------------------
#     # FREEZE HEADER
#     # --------------------------------------------------

#     ws.freeze_panes = "A2"

#     # --------------------------------------------------
#     # ENABLE FILTERS
#     # --------------------------------------------------

#     ws.auto_filter.ref = ws.dimensions

#     # --------------------------------------------------
#     # AUTO COLUMN WIDTH
#     # --------------------------------------------------

#     for col in ws.columns:
#         max_length = max(
#             len(str(cell.value or ""))
#             for cell in col
#         )
#         adjusted_width = min(max_length + 5, 50)
#         ws.column_dimensions[col[0].column_letter].width = adjusted_width

#     # --------------------------------------------------
#     # HIDE COLUMNS
#     # --------------------------------------------------

#     hidden_columns = [
#         "A",  # Budget reference ID
#         "C",  # Partner ID
#         "L",  # Date
#         "M",  # Type of expenses ID
#         "V",  # Declaration
#     ]

#     for col in hidden_columns:
#         ws.column_dimensions[col].hidden = True

#     # --------------------------------------------------
#     # FILE NAME
#     # --------------------------------------------------

#     month = data.get("month") or "month"

#     financial_year = (
#         str(data.get("financial_year") or "fy")
#         .split("-")[0]  # e.g. "2024-2025" → "2024"
#     )

#     filename = f"{month}_{financial_year}_creche_utilisation"

#     # --------------------------------------------------
#     # SAVE FILE
#     # --------------------------------------------------

#     output = BytesIO()

#     wb.save(output)

#     frappe.response["filename"] = f"{filename}.xlsx"
#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "binary"



@frappe.whitelist()
def download_utilisation_template(data=None):

    import json
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    # --------------------------------------------------
    # PARSE DATA
    # --------------------------------------------------

    if isinstance(data, str):
        data = json.loads(data)

    data = data or {}

    # --------------------------------------------------
    # GET LOGGED IN USER
    # --------------------------------------------------

    logged_in_user = frappe.session.user

    # --------------------------------------------------
    # FIND IMPORT TEMPLATE SETTINGS
    # --------------------------------------------------

    template_settings_name = frappe.db.get_value(
        "Import Template settings",
        {
            "user": logged_in_user
        }
    )

    # --------------------------------------------------
    # GET DOCUMENT
    # --------------------------------------------------

    if template_settings_name:

        template_settings = frappe.get_doc(
            "Import Template settings",
            template_settings_name
        )

        items = []

        for row in template_settings.table_jndv:

            items.append({

                "name": row.type_of_expenses_id,

                "budget_main_head": row.budget_main_head,

                "budget_sub_head": row.budget_sub_head,

                "type_of_expenses": row.type_of_expenses
            })

    else:

        items = frappe.get_all(
            "Budget and Expense items list",
            fields=[
                "name",
                "budget_main_head",
                "budget_sub_head",
                "type_of_expenses"
            ],
            order_by="name asc"
        )

    # --------------------------------------------------
    # VALIDATION
    # --------------------------------------------------

    if not items:

        frappe.throw(
            "No Import Template Items found for this user"
        )

    # --------------------------------------------------
    # CREATE WORKBOOK
    # --------------------------------------------------

    wb = Workbook()

    ws = wb.active

    ws.title = "Creche Utilisation"

    # --------------------------------------------------
    # HEADERS
    # --------------------------------------------------

    headers = [
        "Budget reference ID",                            # A
        "Budget reference Name",                          # B
        "Partner ID",                                     # C
        "Partner Name",                                   # D
        "Grant ID",                                       # E
        "Block",                                          # F
        "District",                                       # G
        "State",                                          # H
        "No of creches",                                  # I
        "No of running creches",                          # J
        "Month",                                          # K
        "Financial year",                                 # L
        "Date",                                           # M

        "Type of expenses ID (Utilisation Items List)",   # N
        "Budget main head (Utilisation Items List)",      # O
        "Budget sub head (Utilisation Items List)",       # P
        "Type of expenses (Utilisation Items List)",      # Q

        "Total Amount (Utilisation Items List)",          # R
        "Notes (Utilisation Items List)",                 # S

        "Total Utilisation",                              # T
        "Bank + Cash Balance as at end of month reported",# U
        "Interest from Bank",                             # V
        "Declaration"                                     # W
    ]

    ws.append(headers)

    # --------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------

    header_fill = PatternFill(
        start_color="D9EAF7",
        end_color="D9EAF7",
        fill_type="solid"
    )

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    # --------------------------------------------------
    # PARENT DATA
    # --------------------------------------------------

    parent_data = [
        data.get("budget_reference_id"),        # A
        data.get("budget_reference_name"),      # B
        data.get("partner_id"),                 # C
        data.get("partner_name"),               # D
        data.get("grant_id"),                   # E
        data.get("block"),                      # F
        data.get("district"),                   # G
        data.get("state"),                      # H
        data.get("no_of_creches"),              # I
        data.get("no_of_running_creches") or 0, # J  <-- NEW (default 0)
        data.get("month"),                      # K
        data.get("financial_year"),             # L
        data.get("date")                        # M
    ]

    PARENT_COL_COUNT = len(parent_data)  # now 13

    # --------------------------------------------------
    # ROW RANGE
    # --------------------------------------------------

    start_row = 2
    end_row = start_row + len(items) - 1

    # --------------------------------------------------
    # ADD DATA ROWS
    # --------------------------------------------------

    for idx, item in enumerate(items, start=start_row):

        row = []

        # Parent data only in first row
        row.extend(
            parent_data if idx == start_row else [""] * PARENT_COL_COUNT
        )

        # Child table values
        row.extend([
            item.get("name"),               # N
            item.get("budget_main_head"),   # O
            item.get("budget_sub_head"),    # P
            item.get("type_of_expenses"),   # Q
        ])

        # Editable Total Amount
        row.append(0)                       # R

        # Notes
        row.append("")                      # S

        # Final columns only in first row
        if idx == start_row:
            row.extend([
                f"=SUM(R{start_row}:R{end_row})",  # T - Total Utilisation
                0,                                  # U - Bank + Cash Balance
                0,                                  # V - Interest from Bank
                1                                   # W - Declaration
            ])
        else:
            row.extend(["", "", "", ""])

        ws.append(row)

    # --------------------------------------------------
    # STYLE - TOTAL AMOUNT COLUMN (R) - green tint
    # --------------------------------------------------

    total_fill = PatternFill(
        start_color="E8F4EA",
        end_color="E8F4EA",
        fill_type="solid"
    )

    for row_no in range(2, end_row + 1):
        ws[f"R{row_no}"].fill = total_fill

    # --------------------------------------------------
    # STYLE - TOTAL UTILISATION COLUMN (T) - yellow tint
    # --------------------------------------------------

    budget_fill = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    for row_no in range(start_row, end_row + 1):
        ws[f"T{row_no}"].fill = budget_fill

    # --------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------

    ws.freeze_panes = "A2"

    # --------------------------------------------------
    # ENABLE FILTERS
    # --------------------------------------------------

    ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------
    # AUTO COLUMN WIDTH
    # --------------------------------------------------

    for col in ws.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in col
        )
        adjusted_width = min(max_length + 5, 50)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    # --------------------------------------------------
    # HIDE COLUMNS
    # --------------------------------------------------

    hidden_columns = [
        "A",  # Budget reference ID
        "C",  # Partner ID
        "M",  # Date
        "N",  # Type of expenses ID
        "W",  # Declaration
    ]

    for col in hidden_columns:
        ws.column_dimensions[col].hidden = True

    # --------------------------------------------------
    # FILE NAME
    # --------------------------------------------------

    month = data.get("month") or "month"

    financial_year = (
        str(data.get("financial_year") or "fy")
        .split("-")[0]  # e.g. "2024-2025" → "2024"
    )

    filename = f"{month}_{financial_year}_creche_utilisation"

    # --------------------------------------------------
    # SAVE FILE
    # --------------------------------------------------

    output = BytesIO()

    wb.save(output)

    frappe.response["filename"] = f"{filename}.xlsx"
    frappe.response["filecontent"] = output.getvalue()
    frappe.response["type"] = "binary"